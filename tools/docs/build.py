"""开发记录文档工具链入口：生成配图 → 渲染 Word 与 PPT → 结构检查。

    python tools/docs/build.py                       # 输出到 build/docs
    python tools/docs/build.py --out D:\\记录            # 输出到指定目录
    python tools/docs/build.py --copy D:\\交付            # 顺便把两份文档复制过去
    python tools/docs/build.py --figures-only         # 只重画配图
    python tools/docs/build.py --skip-check           # 跳过结构检查

产物（.docx / .pptx）不进版本库，build/ 已在 .gitignore 里。
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
from pathlib import Path

# 让脚本既能 `python tools/docs/build.py`、也能 `python -m tools.docs.build` 跑
REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from tools.docs import content  # noqa: E402
from tools.docs import figures as figures_module  # noqa: E402
from tools.docs import render_docx, render_pptx  # noqa: E402
from tools.docs.paths import figures_dir, out_dir, resolve_figure  # noqa: E402

CHECK_CANDIDATES = [
    os.environ.get("DSH_OFFICE_CHECK", ""),
    r"C:\Users\27912\AppData\Local\Programs\DeepSeek Harness\resources\runtime"
    r"\office-skills\scripts\check_office.py",
]


def structural_check(targets: list[Path]) -> bool:
    """用捆绑的结构检查脚本验一遍（找不到就跳过，不当失败）。"""

    script = next((Path(path) for path in CHECK_CANDIDATES if path and Path(path).exists()), None)
    if script is None:
        print("结构检查：跳过（未找到 check_office.py，可设 DSH_OFFICE_CHECK 指定）")
        return True
    ok = True
    for target in targets:
        report = target.with_suffix(".checks.json")
        result = subprocess.run([sys.executable, str(script), str(target), "--out", str(report)],
                                capture_output=True, text=True, encoding="utf-8")
        verdict = "pass" if result.returncode == 0 else "FAIL"
        print(f"结构检查 {target.name}: {verdict}")
        if result.returncode != 0:
            ok = False
            print(result.stdout[-800:] or result.stderr[-800:])
    return ok


def validate(figure_dir: Path) -> list[str]:
    """内容源自检：章号顺序、PPT 页序、配图是否存在。

    这些是「以后加内容」最容易犯的错，挡在建文档之前，比渲染出来再肉眼找便宜。
    """

    problems: list[str] = []
    doc_index = 0
    for feature in content.FEATURES:
        key = feature.get("key", "?")
        heading = feature.get("doc_heading")
        if heading:
            doc_index += 1
            expected = f"3.{doc_index} "
            if not heading.startswith(expected):
                problems.append(f"{key}: doc_heading 是 {heading!r}，按顺序应为 {expected!r} 开头")
        for field in ("figure", "figure2"):
            if not feature.get(field):
                continue
            spec = feature[field][0]
            if not Path(resolve_figure(spec, figure_dir)).exists():
                problems.append(f"{key}: {field} 指向的图不存在 —— {spec}")
        deck = feature.get("deck") or {}
        if deck.get("picture") and not Path(resolve_figure(deck["picture"][0], figure_dir)).exists():
            problems.append(f"{key}: deck.picture 指向的图不存在 —— {deck['picture'][0]}")
        if deck.get("table") and deck["table"] not in content.TABLES:
            problems.append(f"{key}: deck.table 引用了不存在的表 {deck['table']!r}")
    keys = {feature["key"] for feature in content.FEATURES}
    for key in content.DECK_ORDER:
        if key not in keys and key not in render_pptx.STRUCTURE:
            problems.append(f"DECK_ORDER 里的 {key!r} 既不是结构页也不在 FEATURES 里")
    for key in keys:
        if key in content.DECK_ORDER:
            feature = next(item for item in content.FEATURES if item["key"] == key)
            if not feature.get("deck"):
                problems.append(f"{key}: 排进了 PPT 页序，却没有 deck 内容")
    return problems


def text_sanity(pptx_path: Path) -> list[str]:
    """PPT 输出自检：有没有异常字号的文本。

    内容源里把 `("要点", True)` 这种简写写错规格时，字号会被当成 True 渲成 1pt——
    看起来就是「这一段不见了」，肉眼极容易漏。
    """

    from pptx import Presentation
    from pptx.util import Pt

    problems: list[str] = []
    presentation = Presentation(pptx_path)
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None and run.font.size < Pt(8):
                        problems.append(
                            f"第 {index} 页有 {run.font.size.pt:g}pt 的文本：{run.text[:24]!r}"
                        )
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description="生成开发记录文档与配图")
    parser.add_argument("--out", default="build/docs", help="产物目录（默认 build/docs）")
    parser.add_argument("--figures", default=None, help="配图目录（默认 <out>/figures）")
    parser.add_argument("--copy", default=None, help="把两份文档复制到这个目录")
    parser.add_argument("--figures-only", action="store_true", help="只重画配图")
    parser.add_argument("--skip-check", action="store_true", help="跳过结构检查")
    args = parser.parse_args()

    out = out_dir(args.out)
    figure_dir = Path(args.figures) if args.figures else figures_dir(args.out)

    print(f"仓库：{REPO}")
    print(f"产物目录：{out}")
    print("1/3 生成配图")
    figures_module.build(figure_dir)
    problems = validate(figure_dir)
    if problems:
        print("内容源自检没通过：")
        for line in problems:
            print(f"  - {line}")
        return 2
    if args.figures_only:
        return 0

    print("2/3 渲染文档")
    docx_path = render_docx.build(out, figure_dir)
    pptx_path = render_pptx.build(out, figure_dir)
    print(f"  {docx_path}")
    print(f"  {pptx_path}")

    ok = True
    if not args.skip_check:
        print("3/3 结构与输出检查")
        ok = structural_check([docx_path, pptx_path])
        for line in text_sanity(pptx_path):
            print(f"  - {line}")
            ok = False

    if args.copy:
        destination = Path(args.copy)
        destination.mkdir(parents=True, exist_ok=True)
        for path in (docx_path, pptx_path):
            shutil.copy2(path, destination / path.name)
            print(f"  已复制 {destination / path.name}")

    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
